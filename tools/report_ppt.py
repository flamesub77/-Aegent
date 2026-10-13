"""AI 보고서 (prd_v4 P1 → prd_v5 R1): files/PPT 양식.pptx로 그래프 위주 2~3장.

v5: 양식 틀·현대하모니 유지 + 첫 장 KPI 타일 띠, 구역 색 부제목 탭, 쪽 번호, 표 스타일(숫자 오른쪽·기준 미달 빨강).
AI 모드면 패널 해석·다음 조치 문장을 AI가 쓴다 (ctx["ai_text"], 숫자는 분석 결과만 넘김). 이모지 대신 색 막대.

양식 = 4:3 슬라이드 1장: 맨 위 제목, 가로선, 좌·우 패널(직사각형) 2개, 패널마다 둥근 "부제목" 탭.
→ 이 슬라이드를 복제해 한 장에 그래프 2개씩 넣는다. 제목 = 그 장의 결론 한 줄, 부제목 = 그래프 이름,
패널 = 그림 + 아래 해석 1~2줄. 그림은 matplotlib + 현대하모니 글꼴 파일로 그린다 (Plotly 그림 저장은 크롬이 필요해
오프라인 PC에서 위험). 미리보기는 같은 그림을 양식 배치대로 합친 PNG. 수치는 분석 결과(ctx)에서만 가져온다.
"""
import io
from copy import deepcopy
from datetime import date

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib import font_manager as fm  # noqa: E402

import config  # noqa: E402
from tools.domain import PROP_SHORT, PROP_UNIT, factor_label, factor_unit  # noqa: E402
from tools.narrative import band_text  # noqa: E402

TEMPLATE = config.DATA_DIR / "PPT 양식.pptx"
FONT_DIR = config.BASE_DIR / "static"
PPT_FONT = {"L": "HDharmony L", "M": "HDharmony M", "B": "HDharmony B"}  # PPT 글자 (받는 PC에 설치된 글꼴 이름)
BLUE, DARK, ORANGE, GRAY, RED = "#3B82F6", "#1E3A8A", "#EA7A2C", "#B6C2D2", "#DC2626"
MIN_ITEMS, MAX_ITEMS = 4, 6  # 2~3장 × 2개
# 구역 색 (prd_v5 R1): 부제목 탭과 미리보기 탭 색
SECTION_COLOR = {"summary": "1E3A8A", "production": "0284C7", "distribution": "1D4ED8", "monthly": "1D4ED8",
                 "thickness": "7C3AED", "spc": "EA7A2C", "ranking": "0F766E", "scatter": "0F766E", "defect": "DC2626",
                 "actions": "15803D"}

# 항목: key → (이름, 기본 체크)
ITEMS = {
    "summary": ("요약", True), "production": ("연도별 생산량", False), "distribution": ("재질 분포", True),
    "monthly": ("월별 추이", False), "thickness": ("두께 그룹별 재질", True), "spc": ("EWMA·CUSUM 조기 감시", True),
    "ranking": ("원인 가능성 순위", True), "scatter": ("1순위 인자 산점도", True), "defect": ("표면결함", False),
    "actions": ("공정제어 체크리스트", False),
}


# ---------------------------------------------------------------------------
# 글꼴
# ---------------------------------------------------------------------------
def _fp(weight="L", size=10) -> fm.FontProperties:
    path = FONT_DIR / f"HDharmony-{weight}.ttf"
    if path.exists():
        return fm.FontProperties(fname=str(path), size=size)
    return fm.FontProperties(family=["Malgun Gothic", "sans-serif"], size=size)  # 글꼴 파일이 없으면 대체


def _style(ax, title=None):
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.grid(True, color="#EEF2F7", linewidth=0.6)
    ax.set_axisbelow(True)
    for lab in ax.get_xticklabels() + ax.get_yticklabels():
        lab.set_fontproperties(_fp("L", 8))
    if title:
        ax.set_title(title, fontproperties=_fp("M", 10), loc="left")


def _labels(ax, x=None, y=None):
    if x:
        ax.set_xlabel(x, fontproperties=_fp("L", 9))
    if y:
        ax.set_ylabel(y, fontproperties=_fp("L", 9))


def _spec(ax, spec, axis="x"):
    for v in spec:
        if v is not None:
            (ax.axvline if axis == "x" else ax.axhline)(v, color=RED, linestyle="--", linewidth=1)


def _png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def _short(prop) -> str:
    """PPT용 재질 이름: 현대하모니에 없는 r̄(결합 문자)는 'r평균'으로."""
    return PROP_SHORT.get(prop, prop).replace("r̄", "r평균")


def _pl(prop):
    u = PROP_UNIT.get(prop, "-")
    return _short(prop) + ("" if u == "-" else f" ({u})")


# ---------------------------------------------------------------------------
# 항목별 그림 (ctx는 app.py가 화면과 같은 분석 결과로 만든다)
# ---------------------------------------------------------------------------
def _fig(w=4.8, h=3.9):
    return plt.subplots(figsize=(w, h))


def _is_num(v) -> bool:
    try:
        float(str(v).rstrip("%"))
        return True
    except ValueError:
        return False


def _table_png(header, rows, col_w=None, bad=None):
    """표: 머리 줄 연한 파랑, 숫자 오른쪽 정렬, bad(행, 열) 칸은 빨강 (prd_v5 R1)."""
    fig, ax = plt.subplots(figsize=(4.8, 0.45 + 0.36 * (len(rows) + 1)))
    ax.axis("off")
    t = ax.table(cellText=rows, colLabels=header, loc="center", cellLoc="center", colWidths=col_w)
    t.scale(1, 1.5)
    bad = bad or set()
    for (r, c), cell in t.get_celld().items():
        cell.set_edgecolor("#DCE3EC")
        cell.set_linewidth(0.6)
        txt = cell.get_text()
        txt.set_fontproperties(_fp("B" if r == 0 else "L", 8.5))
        if r == 0:
            cell.set_facecolor("#E0ECFF")
            txt.set_color("#1E3A8A")
            continue
        cell.set_facecolor("#FFFFFF" if r % 2 else "#F8FAFC")
        if c > 0 and _is_num(txt.get_text()):
            txt.set_horizontalalignment("right")
            cell.PAD = 0.06
        if (r - 1, c) in bad:
            txt.set_color(RED)
            txt.set_fontproperties(_fp("B", 8.5))
    return _png(fig)


def render(key: str, ctx: dict) -> dict | None:
    """항목 하나 → {"subtitle", "png", "title"(결론), "note"[1~2줄]}. 만들 수 없으면 None."""
    fp = ctx["focus_prop"]
    if key == "summary":
        rows = [[_short(q), v["spec"], v["mean"], v["std"], v["rate"]] for q, v in ctx["summary"].items()]
        if not rows:
            return None
        bad = {(i, 4) for i, r in enumerate(rows) if _is_num(r[4]) and float(str(r[4]).rstrip("%")) < 95}
        return {"subtitle": "재질 요약", "png": _table_png(["재질", "SPEC", "평균", "표준편차", "만족률"], rows, bad=bad),
                "title": ctx["headline"], "note": ctx["summary_lines"][:2]}
    if key == "production":
        y = ctx["yearly"]
        if y is None or y.empty:
            return None
        fig, ax = _fig()
        ax.bar(y["year"].astype(str), y["coils"], color=BLUE)
        for x, c in zip(y["year"].astype(str), y["coils"]):
            ax.text(x, c, f"{int(c):,}", ha="center", va="bottom", fontproperties=_fp("M", 8))
        ax2 = ax.twinx()
        ax2.plot(y["year"].astype(str), y["tons"], color=DARK, marker="o")
        ax2.spines["top"].set_visible(False)
        for lab in ax2.get_yticklabels():
            lab.set_fontproperties(_fp("L", 8))
        _style(ax)
        _labels(ax, None, "코일 수")
        ax2.set_ylabel("중량(톤)", fontproperties=_fp("M", 9))
        return {"subtitle": "연도별 생산량", "png": _png(fig), **_note(ctx, "production")}
    if key == "distribution":
        g, r = ctx["dist_group"], ctx["dist_ref"]
        if g.empty:
            return None
        fig, ax = _fig()
        bins = np.histogram_bin_edges(pd.concat([g[fp], r[fp]]).dropna(), bins=30)
        if len(r):
            ax.hist(r[fp].dropna(), bins=bins, color=GRAY, alpha=0.7, density=True, label=f"{ctx['ref_name']} ({len(r)})")
        ax.hist(g[fp].dropna(), bins=bins, color=ctx["group_color"], alpha=0.85, density=True,
                label=f"{ctx['group_name']} ({len(g)})")
        _spec(ax, ctx["spec"][fp])
        _style(ax)
        _labels(ax, _pl(fp), "비율")
        ax.legend(prop=_fp("L", 8), frameon=False)
        return {"subtitle": f"{PROP_SHORT[fp]} 분포", "png": _png(fig), **_note(ctx, "distribution")}
    if key == "monthly":
        m = ctx["coils"].groupby("month")[fp].agg(["mean", "std"]).dropna()
        if m.empty:
            return None
        fig, ax = _fig()
        x = np.arange(len(m))
        ax.fill_between(x, m["mean"] - m["std"], m["mean"] + m["std"], color=BLUE, alpha=0.15, label="평균 ± 표준편차")
        ax.plot(x, m["mean"], color=DARK, marker="o", markersize=3, label="월평균")
        step = max(1, len(m) // 8)
        ax.set_xticks(x[::step], [s[2:] for s in m.index[::step]])
        _spec(ax, ctx["spec"][fp], "y")
        _style(ax)
        _labels(ax, "월", _pl(fp))
        ax.legend(prop=_fp("L", 8), frameon=False)
        return {"subtitle": f"월별 {PROP_SHORT[fp]} 추이", "png": _png(fig), **_note(ctx, "monthly")}
    if key == "thickness":
        tg = ctx["thickness"].get(fp)
        d = ctx["coils"]
        if tg is None or tg.empty:
            return None
        fig, ax = _fig()
        data = [d.loc[ctx["bands"] == b, fp].dropna() for b in tg["band"]]
        bp = ax.boxplot(data, patch_artist=True, showmeans=True)
        worst = tg.loc[tg["spec_rate"].idxmin(), "band"] if tg["spec_rate"].notna().any() else None
        for patch, b in zip(bp["boxes"], tg["band"]):
            patch.set_facecolor(ORANGE if b == worst else "#BFDBFE")
        ax.set_xticks(range(1, len(tg) + 1), [f"{band_text(b)}\n(n={n})" for b, n in zip(tg["band"], tg["n"])])
        _spec(ax, ctx["spec"][fp], "y")
        _style(ax)
        _labels(ax, "두께 구간 (mm)", _pl(fp))
        return {"subtitle": f"두께 그룹별 {PROP_SHORT[fp]}", "png": _png(fig), **_note(ctx, "thickness")}
    if key == "spc":
        res = ctx.get("spc")
        if not res or "error" in res:
            return None
        pts, h = res["points"], res["params"]["h"]
        fig, (a1, a2) = plt.subplots(2, 1, figsize=(4.8, 3.9), sharex=True, gridspec_kw={"height_ratios": [2, 1]})
        a1.scatter(pts["date"], pts[fp], s=3, color=GRAY)
        a1.plot(pts["date"], pts["ewma"], color=DARK, linewidth=1.4, label="EWMA")
        a1.plot(pts["date"], pts["ucl"], color="#64748B", linestyle=":", linewidth=1)
        a1.plot(pts["date"], pts["lcl"], color="#64748B", linestyle=":", linewidth=1)
        _spec(a1, ctx["spec"][fp], "y")
        a2.plot(pts["date"], pts["cusum_hi"], color=ORANGE, linewidth=1, label="CUSUM 높은 쪽")
        a2.plot(pts["date"], pts["cusum_lo"], color=DARK, linewidth=1, label="CUSUM 낮은 쪽")
        a2.axhline(h, color=RED, linestyle=":", linewidth=1)
        alarm = (pts["ewma_alarm"] | pts["cusum_alarm"]).to_numpy()
        for a in (a1, a2):
            a.fill_between(pts["date"], 0, 1, where=alarm, color=ORANGE, alpha=0.12, transform=a.get_xaxis_transform())
            _style(a)
        _labels(a1, None, _pl(fp))
        _labels(a2, None, "CUSUM")
        a1.legend(prop=_fp("L", 7), frameon=False, loc="upper left")
        fig.autofmt_xdate()
        return {"subtitle": f"{PROP_SHORT[fp]} 조기 감시 (EWMA·CUSUM)", "png": _png(fig), **_note(ctx, "spc")}
    if key == "ranking":
        rows = ctx["comp"]["ranking"][:8][::-1]
        if not rows:
            return None
        fig, ax = _fig()
        vals = [r["score"] for r in rows]
        ax.barh([r["label"] for r in rows], vals,
                color=[ORANGE if (r["diff"] or 0) > 0 else DARK for r in rows],
                alpha=1.0)
        for i, (r, v) in enumerate(zip(rows, vals)):
            ax.text(v, i, f" {v:+.2f}" + (" 원리 일치" if r["mechanism"]["status"] == "일치" else ""), va="center",
                    fontproperties=_fp("M", 8))
        ax.axvline(0, color="#94A3B8", linewidth=0.8)
        _style(ax)
        _labels(ax, ctx["score_label"], None)
        return {"subtitle": f"원인 가능성 순위 ({ctx['model_name']})", "png": _png(fig), **_note(ctx, "ranking")}
    if key == "scatter":
        f = ctx["factor"]
        g, r = ctx["group"], ctx["ref"]
        if f not in g.columns:
            return None
        fig, ax = _fig()
        ax.scatter(r[f], r[fp], s=6, color=GRAY, alpha=0.6, label=f"{ctx['ref_name']} ({len(r)})")
        ax.scatter(g[f], g[fp], s=9, color=ctx["group_color"], alpha=0.9, label=f"{ctx['group_name']} ({len(g)})")
        _spec(ax, ctx["spec"][fp], "y")
        _style(ax)
        _labels(ax, f"{factor_label(f)} ({factor_unit(f)})", _pl(fp))
        ax.legend(prop=_fp("L", 8), frameon=False)
        return {"subtitle": f"{factor_label(f)} vs {PROP_SHORT[fp]}", "png": _png(fig), **_note(ctx, "scatter")}
    if key == "defect":
        dg = ctx.get("defect_grades")
        if dg is not None and not dg.empty:
            dg = dg.sort_values("rate_target_pct")
            fig, ax = _fig()
            ax.barh(dg["grade"], dg["rate_baseline_pct"].fillna(0), color=GRAY, label="나머지 기간", height=0.4,
                    align="edge")
            ax.barh(dg["grade"], dg["rate_target_pct"], color=RED, label="이 기간", height=-0.4, align="edge")
            _style(ax)
            _labels(ax, "결함률 (%)", None)
            ax.legend(prop=_fp("L", 8), frameon=False)
            return {"subtitle": f"강종별 {ctx['defect']} 결함률", "png": _png(fig), **_note(ctx, "defect")}
        dm = ctx.get("defect_monthly")
        if dm is None or dm.empty:
            return None
        fig, ax = _fig()
        ax.bar(dm.index.str[2:], dm.values, color=RED)
        step = max(1, len(dm) // 8)
        ax.set_xticks(range(0, len(dm), step), [s[2:] for s in dm.index[::step]])
        _style(ax)
        _labels(ax, "월", "결함률 (%)")
        return {"subtitle": f"월별 {ctx['defect']} 결함률", "png": _png(fig), **_note(ctx, "defect")}
    if key == "actions":
        acts = ((ctx.get("ai_text") or {}).get("actions") or ctx["actions"]
                or ["눈에 띄는 원인 후보가 없어요. 시험 조건부터 확인해 보세요."])
        rows = [[i + 1, a if len(a) <= 38 else a[:37] + "…"] for i, a in enumerate(acts[:6])]
        return {"subtitle": "공정제어 체크리스트", "png": _table_png(["#", "조치"], rows, [0.08, 0.92]),
                "title": "다음 조치", "note": ["원인 순위에서 야금 원리로 설명되는 인자의 조치예요."]}
    return None


def render_item(key: str, ctx: dict) -> dict | None:
    """render + 구역 색 + (AI 모드면) AI가 쓴 해석 문장."""
    p = render(key, ctx)
    if p is None:
        return None
    p["key"], p["color"] = key, SECTION_COLOR.get(key, "1D4ED8")
    ai = (ctx.get("ai_text") or {}).get("notes", {}).get(key)
    if ai:
        p["note"] = [str(x) for x in ai][:2]
    elif key == "summary" and (ctx.get("ai_text") or {}).get("summary"):
        p["note"] = ctx["ai_text"]["summary"][:2]
    return p


def _note(ctx, key):
    title, lines = ctx["notes"].get(key, ("", []))
    return {"title": title, "note": [x.replace("**", "") for x in lines[:2]]}


def available(ctx: dict) -> list[str]:
    """지금 조건으로 만들 수 있는 항목."""
    out = []
    for k in ITEMS:
        if k == "spc" and (not ctx.get("spc") or "error" in ctx["spc"]):
            continue
        if k in ("ranking", "scatter") and not ctx["comp"]["ranking"]:
            continue
        if k == "thickness" and (ctx["thickness"].get(ctx["focus_prop"]) is None
                                 or ctx["thickness"][ctx["focus_prop"]].empty):
            continue
        out.append(k)
    return out


def check_count(n: int) -> str | None:
    if n < MIN_ITEMS:
        return f"2장을 채우려면 {MIN_ITEMS - n}개 이상 더 골라 주세요 (2~3장, 한 장에 그래프 2개)."
    if n > MAX_ITEMS:
        return f"3장까지라 {MAX_ITEMS}개만 고를 수 있어요."
    return None


def pages(keys: list[str], ctx: dict) -> list[dict]:
    """고른 항목 → 장 목록 [{title, panels: [panel, panel]}]. 홀수면 마지막 칸에 요약."""
    panels = [p for p in (render_item(k, ctx) for k in keys) if p]
    if len(panels) % 2 and "summary" not in keys:
        extra = render_item("summary", ctx)
        if extra:
            panels.append(extra)
    out = []
    for i in range(0, len(panels), 2):
        pair = panels[i:i + 2]
        title = ctx["cover_title"] if i == 0 else (pair[0]["title"] or ctx["cover_title"])
        out.append({"title": title, "panels": pair})
    if out:
        out[0]["kpis"] = ctx.get("kpis") or []
    return out


# ---------------------------------------------------------------------------
# PPT 쓰기 (python-pptx)
# ---------------------------------------------------------------------------
def _set_text(shape, text, weight="L", size=None, color=None):
    from pptx.oxml.ns import qn
    from pptx.util import Pt
    tf = shape.text_frame
    p = tf.paragraphs[0]
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    runs = p.runs
    run = runs[0] if runs else p.add_run()
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    run.text = text
    font = run.font
    font.name = PPT_FONT[weight]
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):  # 한글은 동아시아 글꼴 자리에도 넣어야 적용된다
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", PPT_FONT[weight])
    if size:
        font.size = Pt(size)
    if color:
        from pptx.dml.color import RGBColor
        font.color.rgb = RGBColor.from_string(color)


def _layout_shapes(slide, slide_h):
    """양식 도형 찾기: 이름 대신 위치로 (양식이 바뀌어도 동작). 반환 (제목, [부제목 왼·오], [패널 왼·오])."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    shapes = list(slide.shapes)
    texts = [s for s in shapes if s.has_text_frame]
    title = min(texts, key=lambda s: s.top) if texts else None
    rects = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and s is not title]
    subs = sorted([s for s in rects if s.has_text_frame and s.text_frame.text.strip()], key=lambda s: s.left)[:2]
    panels = sorted([s for s in rects if s not in subs and s.height > slide_h * 0.4], key=lambda s: s.left)[:2]
    return title, subs, panels


def _clone(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for ph in list(new.placeholders):  # 레이아웃이 넣은 빈 자리 지우기
        ph._element.getparent().remove(ph._element)
    for shp in src.shapes:
        new.shapes._spTree.insert_element_before(deepcopy(shp._element), "p:extLst")
    return new


KPI_COLOR = {"ok": ("ECFDF5", "15803D"), "bad": ("FEF2F2", "DC2626"), "up": ("FFF7ED", "EA7A2C"),
             "down": ("EFF6FF", "1E3A8A")}


def _font(run, weight, size, color):
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Pt
    run.font.name = PPT_FONT[weight]
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", PPT_FONT[weight])


def _kpi_tiles(slide, kpis: list[dict], left, top, width, height):
    """첫 장 KPI 타일 띠: 재질마다 이름·평균·만족률, 판정 색 (prd_v5 R1)."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu
    kpis = kpis[:4]
    if not kpis:
        return
    gap = int(width * 0.02)
    w = int((width - gap * (len(kpis) - 1)) / len(kpis))
    for i, k in enumerate(kpis):
        fill, line = KPI_COLOR.get(k.get("state", "ok"), KPI_COLOR["ok"])
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(left + i * (w + gap)), Emu(top), Emu(w), Emu(height))
        box.adjustments[0] = 0.18
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string(fill)
        box.line.color.rgb = RGBColor.from_string(line)
        box.line.width = Emu(9525)
        tf = box.text_frame
        tf.margin_left = tf.margin_right = Emu(36000)
        tf.margin_top = tf.margin_bottom = Emu(0)
        _set_text(box, f"{k['name'].replace('r̄', 'r평균')} {k['mean']}", "B", 10, line)
        run = tf.add_paragraph().add_run()
        run.text = k.get("sub", "")
        _font(run, "L", 7, "475569")


def build_pptx(page_list: list[dict], footer: str) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Pt
    prs = Presentation(str(TEMPLATE))
    base = prs.slides[0]
    slides = [base] + [_clone(prs, base) for _ in page_list[1:]]
    sw, sh = prs.slide_width, prs.slide_height
    for n, (slide, page) in enumerate(zip(slides, page_list), start=1):
        title, subs, panels = _layout_shapes(slide, prs.slide_height)
        if title is not None:
            _set_text(title, page["title"], "B")
            if page.get("kpis"):  # 첫 장: 제목은 왼쪽, KPI 타일은 오른쪽
                title.width = int(sw * 0.5)
                _kpi_tiles(slide, page["kpis"], int(sw * 0.53), int(title.top), int(sw * 0.45),
                           int(max(title.height, sh * 0.07)))
        for i, panel_rect in enumerate(panels):
            if i >= len(page["panels"]):
                if i < len(subs):
                    _set_text(subs[i], "", "M")
                continue
            p = page["panels"][i]
            if i < len(subs):
                _set_text(subs[i], p["subtitle"], "M", None, "FFFFFF")
                subs[i].fill.solid()
                subs[i].fill.fore_color.rgb = RGBColor.from_string(p.get("color", "1D4ED8"))
            sub_bottom = subs[i].top + subs[i].height if i < len(subs) else panel_rect.top
            pad = Emu(int(panel_rect.width * 0.04))
            note_h = Emu(int(panel_rect.height * 0.16))
            box_top = max(panel_rect.top, sub_bottom) + pad
            box_w = panel_rect.width - 2 * pad
            box_h = panel_rect.top + panel_rect.height - box_top - note_h - pad
            pic = slide.shapes.add_picture(io.BytesIO(p["png"]), panel_rect.left + pad, box_top, width=box_w)
            if pic.height > box_h:  # 패널 밖으로 넘치지 않게 높이 기준으로 다시 맞춤
                ratio = box_h / pic.height
                pic.width, pic.height = int(pic.width * ratio), int(box_h)
                pic.left = int(panel_rect.left + (panel_rect.width - pic.width) / 2)
            tb = slide.shapes.add_textbox(panel_rect.left + pad, panel_rect.top + panel_rect.height - note_h - pad // 2,
                                          box_w, note_h)
            tb.text_frame.word_wrap = True
            _set_text(tb, " ".join(p["note"]), "L", 9, "334155")
        foot = slide.shapes.add_textbox(Emu(int(sw * 0.02)), Emu(int(sh * 0.955)), Emu(int(sw * 0.86)),
                                        Emu(int(sh * 0.04)))
        _set_text(foot, footer, "L", 7, "64748B")
        foot.text_frame.paragraphs[0].runs[0].font.size = Pt(7)
        num = slide.shapes.add_textbox(Emu(int(sw * 0.88)), Emu(int(sh * 0.955)), Emu(int(sw * 0.10)), Emu(int(sh * 0.04)))
        _set_text(num, f"{n} / {len(page_list)}", "M", 8, "1E3A8A")
        num.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 미리보기 (양식 배치대로 합친 PNG)
# ---------------------------------------------------------------------------
def preview_pngs(page_list: list[dict], footer: str) -> list[bytes]:
    """양식과 같은 배치(4:3, 제목·선·좌우 패널·부제목 탭)로 슬라이드 그림을 만든다."""
    import matplotlib.image as mpimg
    from matplotlib.patches import FancyBboxPatch, Rectangle
    out = []
    W, H = 25.4, 19.05  # cm
    for n, page in enumerate(page_list, start=1):
        fig = plt.figure(figsize=(8, 6))
        ax = fig.add_axes([0, 0, 1, 1])
        ax.set_xlim(0, W)
        ax.set_ylim(H, 0)
        ax.axis("off")
        kp = (page.get("kpis") or [])[:4]
        ax.text(1.5, 1.25, page["title"], fontproperties=_fp("B", 12 if kp else 14), va="center")
        for j, k in enumerate(kp):  # KPI 타일 띠
            fill, line = KPI_COLOR.get(k.get("state", "ok"), KPI_COLOR["ok"])
            w = (11.4 - 0.25 * (len(kp) - 1)) / len(kp)
            x = 13.5 + j * (w + 0.25)
            ax.add_patch(FancyBboxPatch((x, 0.35), w, 1.45, boxstyle="round,pad=0,rounding_size=0.25",
                                        facecolor="#" + fill, edgecolor="#" + line, linewidth=0.8))
            ax.text(x + 0.2, 0.85, f"{k['name'].replace('r̄', 'r평균')} {k['mean']}", fontproperties=_fp("B", 7.5), color="#" + line, va="center")
            ax.text(x + 0.2, 1.4, k.get("sub", ""), fontproperties=_fp("L", 5.5), color="#475569", va="center")
        ax.plot([2.1, 25.2], [2.1, 2.1], color="#1E3A8A", linewidth=1.2)
        for i, (x0, sx) in enumerate(((0.2, 2.8), (12.8, 15.4))):
            ax.add_patch(Rectangle((x0, 3.1), 12.4, 15.5, fill=False, edgecolor="#CBD5E1"))
            if i < len(page["panels"]):
                p = page["panels"][i]
                ax.add_patch(FancyBboxPatch((sx, 2.6), 7.2, 1.1, boxstyle="round,pad=0,rounding_size=0.5",
                                            facecolor="#" + p.get("color", "1D4ED8"), edgecolor="none"))
                ax.text(sx + 3.6, 3.15, p["subtitle"], fontproperties=_fp("M", 8), color="white", ha="center", va="center")
                img = mpimg.imread(io.BytesIO(p["png"]), format="png")
                ih, iw = img.shape[:2]
                bw, bh = 11.6, 11.6
                scale = min(bw / iw, bh / ih)
                dw, dh = iw * scale, ih * scale
                left, top = x0 + (12.4 - dw) / 2, 4.0
                ax.imshow(img, extent=(left, left + dw, top + dh, top), zorder=2)
                ax.text(x0 + 0.5, 16.4, "\n".join(p["note"])[:140], fontproperties=_fp("L", 6.5), color="#334155",
                        va="top", wrap=True)
        ax.text(0.5, 18.85, footer, fontproperties=_fp("L", 5.5), color="#64748B")
        ax.text(24.9, 18.85, f"{n} / {len(page_list)}", fontproperties=_fp("M", 6.5), color="#1E3A8A", ha="right")
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=110, facecolor="white")
        plt.close(fig)
        out.append(buf.getvalue())
    return out


def filename(grade: str) -> str:
    return f"{grade}_AI보고서_{date.today():%Y%m%d}.pptx"


# ---------------------------------------------------------------------------
# AI 모드: 문장만 AI가 쓴다 (prd_v5 R1). 숫자는 ctx의 분석 결과만 넘긴다 (원본 행은 보내지 않음)
# ---------------------------------------------------------------------------
AI_SYSTEM = ("당신은 냉연강판 품질 보고서를 쓰는 연구개발자입니다. 주어진 숫자만 쓰고 새 숫자를 만들지 마세요. "
             "한국어, 짧게, 숫자 먼저, 화살표(↑↓→) 사용. 반드시 JSON 하나만 출력하세요.")


def ai_prompt(keys: list[str], ctx: dict) -> str:
    import json
    items = {}
    for k in keys:
        if k in ctx["notes"]:
            t, lines = ctx["notes"][k]
            items[k] = {"그래프": ITEMS[k][0], "결론": t, "해석": lines}
    data = {"강종": ctx["grade"], "조건": ctx["footer"], "재질 요약": ctx["summary"], "항목": items,
            "원인 순위": [{"인자": r["label"], "차이": r["diff"], "단위": r["unit"], "원리": r["mechanism"]["status"]}
                      for r in ctx["comp"]["ranking"][:3]],
            "규칙 조치": ctx["actions"][:4]}
    return ("아래 분석 결과로 AI 보고서 문장을 써 주세요.\n"
            '형식: {"summary": ["요약 2~3줄"], "notes": {"항목키": ["해석 1~2줄 (40자 안팎)"]}, "actions": ["다음 조치 1~4개"]}\n'
            "notes의 키는 '항목'의 키만 쓰세요.\n\n" + json.dumps(data, ensure_ascii=False, default=str))


def parse_ai(text: str) -> dict | None:
    """AI 응답(JSON) → {"summary", "notes", "actions"}. 형식이 틀리면 None (규칙 문장 유지)."""
    import json
    import re
    m = re.search(r"\{.*\}", text or "", re.S)
    if not m:
        return None
    try:
        d = json.loads(m.group(0))
    except json.JSONDecodeError:
        return None
    if not isinstance(d, dict):
        return None
    notes = d.get("notes") if isinstance(d.get("notes"), dict) else {}
    return {"summary": [str(x) for x in d.get("summary") or []][:3],
            "notes": {k: v if isinstance(v, list) else [str(v)] for k, v in notes.items() if k in ITEMS},
            "actions": [str(x) for x in d.get("actions") or []][:6]}
