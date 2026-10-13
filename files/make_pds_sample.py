"""회사 양식(coil data(pds).xlsx)으로 샘플 가상 데이터를 만든다 (prd_v3 D1, prd_v5 D1).

- 2026년 1~6월 기존 6강종: 기존 샘플(coil_data.csv) 4,000코일 값을 그대로 옮긴다 → S1~S6 정답 유지.
- 2019~2025년: 원래 생성기(generate_data.py)의 공정·재질 모델로 새로 만든다 (13강종, 이상 시나리오 없음).
  평가와 무관한 작은 흐름 변화(DRIFTS)를 몇 개 넣어 조기 감시·비교 화면을 써 볼 수 있게 한다 (정답지에 넣지 않음).
- 2026년 1~6월 신규 7강종: 따로 만들어 붙인다. S6(GA재 파우더링)이 흔들리지 않게 CGL_GA 라인에는 넣지 않는다.
- 신규 강종(DP·TRIP·MART 등)은 생성기의 계열식(HSLA 등)을 빌려 계산하고, 화면용 계열은 files/grade_family.csv에 적는다.
- 라인: CAL·CGL_GA + EGL(전기도금, CAL재 일부) + CGL_GI(용융아연, CGL_GA재 일부, 합금화 없음). 공장 3곳.
- 결함: 생성기 5종 + 롤 마크·스크래치 (조질 연신·라인속도에 따라).
- 양식 165열은 그대로 두고, 양식에 없는 항목은 맨 뒤 "추가 열"에 쓴다 (tools/pds.py EXTRA).

데이터 생성 스크립트다. 앱·에이전트·도구 코드는 이 파일과 generate_data.py를 읽지 않는다.
실행: python files/make_pds_sample.py [저장 경로, 기본은 양식 파일 자체]
"""
import importlib.util
from datetime import timedelta
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
XLSX = HERE / "coil data(pds).xlsx"
SEED = 20261012
HISTORY = {2019: 3000, 2020: 3100, 2021: 3300, 2022: 3500, 2023: 3700, 2024: 3900, 2025: 4200}
NEW_2026 = 1100  # 2026년 1~6월 신규 강종 코일 수

# 신규 강종: 생성기 GRADES 형식 (calc = 재질 계산에 빌려 쓸 계열, family = 화면용 계열)
NEW_GRADES = {
    "IF260_IN": dict(name="극저탄 내판 260급", family="IF", calc="IF", ys=(110, 180), ts=(260, 340), el_min=46, bh_min=None,
                     chem=dict(C=0.0018, Mn=0.08, Si=0.010, P=0.008, S=0.006, N=0.0024, Al=0.035, Ti=0.055, Nb=0.000, B=0.0),
                     std=dict(SRT=1200, FDT=920, CT=710, RED=82, SS=830, OA=None, GA=520, SPM=0.8),
                     base=dict(YS=140, TS=286, EL=50.5), lines={"CAL": 0.5, "CGL_GA": 0.5}, share=0.08),
    "BH180_OUT": dict(name="BH 외판 180급", family="BH", calc="BH", ys=(180, 240), ts=(300, 380), el_min=38, bh_min=30,
                      chem=dict(C=0.0022, Mn=0.25, Si=0.015, P=0.020, S=0.006, N=0.0025, Al=0.035, Ti=0.0, Nb=0.007, B=0.0004),
                      std=dict(SRT=1200, FDT=910, CT=650, RED=77, SS=840, OA=None, GA=520, SPM=1.0),
                      base=dict(YS=205, TS=338, EL=40.5), lines={"CAL": 0.4, "CGL_GA": 0.6}, share=0.07),
    "HSLA440": dict(name="석출경화강 440급", family="HSLA", calc="HSLA", ys=(305, 420), ts=(440, 560), el_min=26, bh_min=None,
                    chem=dict(C=0.060, Mn=1.00, Si=0.200, P=0.015, S=0.004, N=0.0040, Al=0.040, Ti=0.0, Nb=0.020, B=0.0),
                    std=dict(SRT=1220, FDT=880, CT=600, RED=62, SS=800, OA=400, GA=520, SPM=1.0),
                    base=dict(YS=362, TS=495, EL=30.0), lines={"CAL": 0.7, "CGL_GA": 0.3}, share=0.07),
    "DP590": dict(name="복합조직(DP)강 590급", family="DP", calc="HSLA", ys=(340, 460), ts=(590, 700), el_min=21, bh_min=None,
                  chem=dict(C=0.080, Mn=1.80, Si=0.300, P=0.012, S=0.003, N=0.0040, Al=0.040, Ti=0.0, Nb=0.0, B=0.0),
                  std=dict(SRT=1220, FDT=880, CT=620, RED=58, SS=810, OA=400, GA=520, SPM=0.4),
                  base=dict(YS=398, TS=640, EL=24.5), lines={"CAL": 0.6, "CGL_GA": 0.4}, share=0.08),
    "DP780": dict(name="복합조직(DP)강 780급", family="DP", calc="HSLA", ys=(440, 600), ts=(780, 900), el_min=14, bh_min=None,
                  chem=dict(C=0.120, Mn=2.10, Si=0.400, P=0.012, S=0.003, N=0.0040, Al=0.040, Ti=0.0, Nb=0.0, B=0.0),
                  std=dict(SRT=1230, FDT=880, CT=600, RED=55, SS=820, OA=400, GA=520, SPM=0.3),
                  base=dict(YS=518, TS=835, EL=17.0), lines={"CAL": 0.7, "CGL_GA": 0.3}, share=0.06),
    "TRIP780": dict(name="변태유기소성(TRIP)강 780급", family="TRIP", calc="HSLA", ys=(440, 600), ts=(780, 900), el_min=21,
                    bh_min=None,
                    chem=dict(C=0.180, Mn=1.60, Si=1.500, P=0.012, S=0.003, N=0.0040, Al=0.040, Ti=0.0, Nb=0.0, B=0.0),
                    std=dict(SRT=1230, FDT=880, CT=580, RED=55, SS=820, OA=420, GA=520, SPM=0.3),
                    base=dict(YS=505, TS=828, EL=24.5), lines={"CAL": 1.0}, share=0.04),
    "MART1180": dict(name="마르텐사이트(MART)강 1180급", family="MART", calc="HSLA", ys=(900, 1150), ts=(1180, 1350),
                     el_min=4, bh_min=None,
                     chem=dict(C=0.220, Mn=1.90, Si=0.300, P=0.010, S=0.002, N=0.0040, Al=0.040, Ti=0.0, Nb=0.0, B=0.0020),
                     std=dict(SRT=1240, FDT=880, CT=560, RED=50, SS=840, OA=None, GA=520, SPM=0.2),
                     base=dict(YS=1020, TS=1265, EL=7.0), lines={"CAL": 1.0}, share=0.02),
}
FORM_BASE_NEW = {"IF260_IN": dict(rbar=2.05, dr=0.45, a=0.25, n=0.245), "BH180_OUT": dict(rbar=1.75, dr=0.35, a=0.20, n=0.225),
                 "HSLA440": dict(rbar=1.10, dr=0.25, a=0.12, n=0.170), "DP590": dict(rbar=0.95, dr=0.20, a=0.10, n=0.170),
                 "DP780": dict(rbar=0.85, dr=0.18, a=0.10, n=0.140), "TRIP780": dict(rbar=0.90, dr=0.18, a=0.10, n=0.200),
                 "MART1180": dict(rbar=0.80, dr=0.15, a=0.08, n=0.060)}
FORM_SPEC_NEW = {"IF260_IN": dict(rbar_min=1.8, n_min=0.23, dr_max=None), "BH180_OUT": dict(rbar_min=1.5, n_min=0.2, dr_max=None)}

ROLL = {"IF270_IN": "CRIF27", "IF340_OUT": "CRIF34", "BH340_OUT": "CRBH34", "LC340": "CR34", "LC440R": "CR44R",
        "HSLA590C": "CR590HS", "IF260_IN": "CRIF26", "BH180_OUT": "CRBH18", "HSLA440": "CR440HS", "DP590": "CR590DP",
        "DP780": "CR780DP", "TRIP780": "CR780TR", "MART1180": "CR1180M"}
AIM = {"IF270_IN": "IF27A", "IF340_OUT": "IF34P", "BH340_OUT": "BH34N", "LC340": "LC34A", "LC440R": "LC44M",
       "HSLA590C": "HS59N", "IF260_IN": "IF26A", "BH180_OUT": "BH18N", "HSLA440": "HS44N", "DP590": "DP59A",
       "DP780": "DP78A", "TRIP780": "TR78S", "MART1180": "MT18B"}
USE = {  # 용도코드 (가정). files/use_codes.csv에 이름
    ("IF270_IN", None): "AI0", ("IF340_OUT", None): "AT0", ("BH340_OUT", "CGL_GA"): "AAT0",
    ("BH340_OUT", "CAL"): "AT0", ("BH340_OUT", None): "AT0", ("LC340", None): "GP0", ("LC440R", None): "ST0", ("HSLA590C", None): "AS0",
    ("IF260_IN", None): "AI0", ("BH180_OUT", None): "AT0", ("BH180_OUT", "CGL_GA"): "AAT0", ("HSLA440", None): "AS0",
    ("DP590", None): "AS0", ("DP780", None): "AR0", ("TRIP780", None): "AR0", ("MART1180", None): "AR1",
}
USE_NAMES = {"AI0": "자동차 내판", "AT0": "자동차 외판", "AAT0": "자동차 외판 (GA)", "AS0": "자동차 구조",
             "AR0": "자동차 보강재", "AR1": "자동차 충돌 보강", "GP0": "일반 가공", "ST0": "구조용"}
PLANT = {"CAL": "1냉연", "EGL": "1냉연", "CGL_GA": "2냉연", "CGL_GI": "3냉연"}
# 평가와 무관한 작은 흐름 변화 (년-월, 강종, 인자, 변화량)
DRIFTS = [("2020-09", "HSLA590C", "CT", 22), ("2021-07", "IF340_OUT", "SS", -18), ("2022-03", "DP590", "SPM", 0.5),
          ("2022-11", "LC440R", "P", 0.012), ("2023-05", "DP780", "SS", -20), ("2024-02", "BH340_OUT", "CT", 20),
          ("2025-08", "HSLA440", "CT", 25), ("2025-04", "IF270_IN", "cast_speed", 0.2)]
EXTRA_HEADERS = ["합금화온도", "도금욕온도", "디스케일링압력", "산세속도", "HCl농도", "냉연결함", "냉연결함등급", "판정",
                 "시생산여부"]
HOT_DEFECTS = {"스케일"}  # 열연 블록 결함 칸에 적을 결함 (나머지는 냉연 결함 추가 열에만)


def _generator():
    spec = importlib.util.spec_from_file_location("generate_data", HERE / "generate_data.py")
    gd = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(gd)
    for code, g in NEW_GRADES.items():  # 신규 강종: 재질 계산은 빌린 계열식(calc)으로
        gd.GRADES[code] = {**g, "family": g["calc"]}
        gd.FORM_BASE[code] = FORM_BASE_NEW[code]
    return gd


def _relabel_lines(df: pd.DataFrame, rng, no_ga: bool = False) -> pd.DataFrame:
    """CAL 일부 → EGL, CGL_GA 일부 → CGL_GI(합금화 없음). no_ga면 CGL_GA를 모두 CGL_GI로."""
    cal = df.index[df["line"] == "CAL"]
    ga = df.index[df["line"] == "CGL_GA"]
    df.loc[rng.choice(cal, size=int(len(cal) * 0.2), replace=False), "line"] = "EGL"
    gi = ga if no_ga else rng.choice(ga, size=int(len(ga) * 0.3), replace=False)
    df.loc[gi, "line"] = "CGL_GI"
    df.loc[df["line"] == "CGL_GI", "GA_temp"] = np.nan
    return df


def _drifts(df: pd.DataFrame) -> pd.DataFrame:
    month = df["prod_date"].dt.strftime("%Y-%m")
    for ym, grade, col, delta in DRIFTS:
        m = (month == ym) & (df["grade"] == grade)
        df.loc[m, col] = df.loc[m, col] + delta
    return df


def _extra_defects(df: pd.DataFrame, rng) -> pd.DataFrame:
    """롤 마크(조질 연신↑), 스크래치(라인속도↑): 결함이 없는 코일에만."""
    sig = lambda x: 1 / (1 + np.exp(-x))  # noqa: E731
    free = df["defect_type"] == "없음"
    p_roll = sig(-5.6 + 2.2 * (df["SPM"] - 1.0))
    p_scr = sig(-6.0 + 0.012 * (df["line_speed"] - 200))
    roll = free & (rng.random(len(df)) < p_roll)
    scr = free & ~roll & (rng.random(len(df)) < p_scr)
    df.loc[roll, "defect_type"], df.loc[scr, "defect_type"] = "롤 마크", "스크래치"
    sev = rng.choice(["경", "중", "중대"], size=len(df), p=[.65, .28, .07])
    df.loc[roll | scr, "defect_grade"] = sev[(roll | scr).to_numpy()]
    return df


def _make(gd, n, start, end, seed, grades=None, no_ga=False) -> pd.DataFrame:
    gd.RNG, gd.RNG2 = np.random.default_rng(seed), np.random.default_rng(seed + 1000)
    rng = np.random.default_rng(seed + 2000)
    saved = dict(gd.GRADES)
    if grades:
        gd.GRADES.clear()
        gd.GRADES.update({k: saved[k] for k in grades})
    try:
        gd.N_COILS, gd.START, gd.END = n, start, end
        df = gd.make_coils()
        df["scenario"] = ""
        df = _relabel_lines(_drifts(df), rng, no_ga)
        df = gd.properties(df)
        df = gd.defects(df)
        df = _extra_defects(df, rng)
        df = gd.lineage(gd.formability(gd.judge(df)))
        df["trial_flag"] = "N"
        for g, idx in df.groupby("grade").groups.items():  # 강종별 시생산 몇 개
            df.loc[rng.choice(idx, size=min(3, len(idx)), replace=False), "trial_flag"] = "Y"
        df["prod_date"] = df["prod_date"].dt.strftime("%Y-%m-%d")
    finally:
        gd.GRADES.clear()
        gd.GRADES.update(saved)
    return df.drop(columns="scenario")


def _offset_ids(df: pd.DataFrame, k: int) -> pd.DataFrame:
    """2026년 신규 강종 코일 번호가 기존 2026년 번호와 겹치지 않게 앞자리를 바꾼다."""
    for col, pre in (("coil_no", "C"), ("heat_no", "H"), ("slab_no", "SL"), ("hot_coil_no", "HC")):
        df[col] = df[col].str.replace(rf"^{pre}(\d{{4}})", rf"{pre}\g<1>{k}", regex=True)
    return df


def history() -> pd.DataFrame:
    """2019~2025년 + 2026년 신규 강종."""
    gd = _generator()
    frames = [_make(gd, n, f"{y}-01-01", f"{y}-12-31", y) for y, n in HISTORY.items()]
    new26 = _make(gd, NEW_2026, "2026-01-01", "2026-06-29", 2026, grades=list(NEW_GRADES), no_ga=True)
    frames.append(_offset_ids(new26, 9))
    return pd.concat(frames, ignore_index=True)


def spec_table(gd=None) -> pd.DataFrame:
    """기존 grade_master + 신규 강종 스펙·표준조건 (양식의 설계·지시 열을 채우는 데 쓴다)."""
    base = pd.read_csv(HERE / "grade_master.csv")
    rows = []
    for code, g in NEW_GRADES.items():
        row = dict(grade=code, grade_name=g["name"], family=g["family"], ys_min=g["ys"][0], ys_max=g["ys"][1],
                   ts_min=g["ts"][0], ts_max=g["ts"][1], el_min=g["el_min"], bh_min=g["bh_min"], lines="/".join(g["lines"]))
        row.update({f"aim_{k}": v for k, v in g["chem"].items()})
        row.update({f"std_{k}": v for k, v in g["std"].items()})
        row.update(FORM_SPEC_NEW.get(code, dict(rbar_min=None, n_min=None, dr_max=None)))
        rows.append(row)
    return pd.concat([base, pd.DataFrame(rows)], ignore_index=True).set_index("grade")


def build(out: Path = XLSX) -> int:
    rng = np.random.default_rng(SEED)
    spec = spec_table()
    now = pd.read_csv(HERE / "coil_data.csv")
    now = now.assign(line=now["line"])  # 기존 2026년 6강종은 값·라인을 그대로 둔다 (S1~S6)
    df = pd.concat([history()[now.columns], now], ignore_index=True).sort_values(["prod_date", "coil_no"])
    df = df.drop_duplicates("coil_no", keep="last")

    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    ws.delete_rows(2, ws.max_row)
    n_template = next(c.column for c in ws[1] if c.value == "BH" and c.column > 150)  # 냉연 블록 끝 (FI)
    ws.delete_cols(n_template + 1, ws.max_column)
    for k, h in enumerate(EXTRA_HEADERS, start=n_template + 1):  # 추가 열 머리글 (회색)
        cell = ws.cell(row=1, column=k, value=h)
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="FFE5E7EB")

    def u(lo, hi, nd=1):
        return round(float(rng.uniform(lo, hi)), nd)

    def nz(v, nd=4):
        return None if pd.isna(v) else round(float(v), nd)

    for _, r in df.iterrows():
        g = spec.loc[r["grade"]]
        roll, aim = ROLL[r["grade"]], AIM[r["grade"]]
        use = USE.get((r["grade"], r["line"])) or USE[(r["grade"], None)]
        cold_day = pd.Timestamp(r["prod_date"])
        hot_day = cold_day - timedelta(days=int(rng.integers(3, 8)))
        cast_day = hot_day - timedelta(days=int(rng.integers(2, 6)))
        steel_day = cast_day - timedelta(hours=int(rng.integers(1, 5)))
        heat, slab, hot, cold = r["heat_no"], r["slab_no"], r["hot_coil_no"], r["coil_no"]
        cast = "CS" + heat[1:7] + f"{int(heat[-2:]) % 10:02d}"
        order, row_no = f"OR{cold_day:%y%m}{rng.integers(1000, 9999)}", int(rng.integers(1, 20))
        tundish = round(1530 + r["superheat"], 1)
        is_if = g["family"] in ("IF", "IF_HS", "BH")
        spec_roll = roll + ("A" if r["thk"] < 1.0 else "B")
        spec_fin = roll + ("G" if r["line"] == "CGL_GA" else "C")
        sn, v, cr, mo, ni, cu = (u(0, 0.003, 4), u(0, 0.004, 4), u(0.01, 0.03, 3),
                                 u(0, 0.005, 4), u(0.005, 0.02, 3), u(0.005, 0.02, 3))
        ceq = r["C"] + r["Mn"] / 6 + (cr + mo + v) / 5 + (ni + cu) / 15
        slab_w = int(r["width"] + 50)
        slab_t = round(0.25 * slab_w / 1000 * 10 * 7.85, 2)
        hot_weight = round(slab_t * 0.985, 2)
        defect = r["defect_type"] if r["defect_type"] != "없음" else ""
        rbar = (r["r0"] + 2 * r["r45"] + r["r90"]) / 4
        hot_ys = {"IF": 200, "IF_HS": 260, "BH": 265, "LC": 300, "HSLA": 490, "DP": 420, "TRIP": 520,
                  "MART": 700}.get(g["family"], 300)

        steelmaking = [
            heat, cast, "냉연", "1제강", steel_day.to_pydatetime(), aim, r["grade"], "RH" if is_if else "LF",
            tundish, int(rng.integers(900_000, 1_400_000)),
            nz(r["C"]), nz(r["Mn"], 2), nz(r["Si"], 3), nz(r["P"]), nz(r["S"]), sn,
            nz(r["Al"]), round(r["Al"] + u(0.002, 0.006, 4), 4), nz(r["Nb"]), nz(r["Ti"]), v, cr, mo, ni, cu,
            nz(r["B"]), u(0, 0.0015, 4), nz(r["N"]), u(1.0, 2.5, 1), round(ceq, 3),
            u(1640, 1690), int(rng.integers(400, 900) if is_if else rng.integers(300, 600)), int(rng.integers(1, 3)),
            int(rng.integers(800, 1600)), int(rng.integers(150, 450)), u(3, 8), None if is_if else u(25, 50),
            None if is_if else nz(r["S"] + u(0.002, 0.006, 4)), None if is_if else int(rng.integers(3000, 9000)),
            None if is_if else u(5, 12), None if is_if else u(3, 8), u(25, 40) if is_if else None,
            f"{u(1590, 1620, 0):.0f}/{u(1565, 1590, 0):.0f}", u(0.5, 1.5, 2) if is_if else None,
            u(10, 20) if is_if else None, int(rng.integers(4000, 9000)), int(rng.integers(0, 3000)),
            u(5, 10), None if is_if else u(5, 10),
            round(g["aim_C"] * 0.8, 4), round(g["aim_C"] * 1.2, 4),
        ]
        casting = [
            slab, heat, cast, "냉연", "1제강", f"{rng.integers(1, 3)}CC", int(rng.integers(1, 3)),
            cast_day.to_pydatetime(), r["grade"], aim, 250, slab_w, 250, slab_w, slab_t, nz(r["cast_speed"], 2),
            u(0.01, 0.06, 2), "MP-IF" if is_if else "MP-LC", "Soft" if is_if else "Medium", f"P{rng.integers(1, 6)}",
            u(92, 99), "정상", "Bifurcated", tundish, "", "없음", "", "", "", "", "", order, row_no,
        ]
        hot_rolling = [
            hot, slab, heat, hot_day.to_pydatetime(), "2열연", nz(r["hot_thk"], 2), int(r["width"] + 10), hot_weight,
            r["grade"], spec_roll, spec_fin, order, aim, use, row_no, g["std_SRT"], 1080, g["std_FDT"], g["std_CT"],
            "U" if is_if else "E", nz(r["SRT"], 1), u(1060, 1100), nz(r["FDT"], 1), nz(r["CT"], 1),
            "U" if is_if else "E", "열연-산세", f"T{hot[-6:]}", "T", "L", None, None, None,
            round(hot_ys + rng.normal(0, 12)), round(hot_ys * 1.25 + rng.normal(0, 12)),
            round({"IF": 45, "IF_HS": 38, "BH": 37, "LC": 33, "HSLA": 24, "DP": 24, "TRIP": 26, "MART": 15}
                  .get(g["family"], 33) + rng.normal(0, 1.5), 1),
            "Y" if defect in HOT_DEFECTS else "N", defect if defect in HOT_DEFECTS else "",
            "D01" if defect in HOT_DEFECTS else "",
        ]
        cold_rolling = [
            cold, hot, hot, slab, heat, cold_day.to_pydatetime(), PLANT.get(r["line"], "1냉연"), r["line"], nz(r["thk"], 2),
            int(r["width"]), round(hot_weight * u(0.45, 0.97, 2), 2), r["grade"], spec_roll, aim, use,
            order, row_no, nz(r["line_speed"], 1), u(700, 760), nz(r["SS"], 1), u(640, 680), u(450, 500),
            u(380, 420) if r["line"] == "CAL" else None, nz(r["OA"], 1), nz(r["SPM"], 2), r["coil_pos"],
            "L", "JIS5", "JIS 5호", f"{g['ys_min']:.0f}~{g['ys_max']:.0f}", f"{g['ts_min']:.0f}~{g['ts_max']:.0f}",
            f"{g['el_min']:.0f}↑", None if pd.isna(g["n_min"]) else f"{g['n_min']}↑",
            None if pd.isna(g["bh_min"]) else f"{g['bh_min']:.0f}↑",
            nz(r["YS"], 0), nz(r["TS"], 0), nz(r["EL"], 1), nz(r["n_value"], 3), nz(r["r0"], 2), nz(r["r45"], 2),
            nz(r["r90"], 2), round(rbar, 3), nz(r["BH"], 0),
        ]
        extra = [nz(r["GA_temp"], 1), nz(r["pot_temp"], 1), nz(r["descale_press"], 1), nz(r["pickle_speed"], 1),
                 nz(r["hcl_conc"], 1), r["defect_type"], None if pd.isna(r["defect_grade"]) else r["defect_grade"],
                 r["judge"], r["trial_flag"]]
        row = steelmaking + casting + hot_rolling + cold_rolling
        assert len(row) == n_template, (len(row), n_template)
        ws.append(row + extra)

    for col in ("E", "BG", "CJ"):
        for cell in ws[col][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm"
    for cell in ws["DX"][1:]:
        cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    wb.save(out)
    pd.DataFrame([{"code": k, "name": v} for k, v in USE_NAMES.items()]).to_csv(
        HERE / "use_codes.csv", index=False, encoding="utf-8-sig")
    spec.reset_index()[["grade", "family", "grade_name"]].to_csv(  # 화면용 강종 계열 (신규 강종 포함)
        HERE / "grade_family.csv", index=False, encoding="utf-8-sig")
    return len(df)


if __name__ == "__main__":
    import sys
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else XLSX
    print(f"{build(out):,}행 작성: {out}")
