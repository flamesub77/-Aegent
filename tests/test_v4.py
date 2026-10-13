"""prd_v4 기능 테스트 (API 없이)."""
import io

import pandas as pd
import pytest

from tools import analysis as A
from tools import dashboard_state as ds
from tools import narrative as N
from tools import spc
from tools.datasource import use_source


@pytest.fixture(autouse=True)
def _sample_source():
    use_source(None)
    yield
    use_source(None)


# M3 조기 감시: 시나리오 달에 방향이 맞는 첫 경보, 이상 없는 달은 경보 없음
@pytest.mark.parametrize("grade,prop,line,start,end,direction", [
    ("HSLA590C", "YS", "CAL", "2026-05-01", "2026-05-31", "낮게"),   # S1
    ("IF340_OUT", "EL", None, "2026-03-01", "2026-03-31", "낮게"),   # S2
    ("LC440R", "TS", None, "2026-02-01", "2026-02-28", "높게"),      # S3
    ("LC340", "YS", None, "2026-04-01", "2026-04-30", "높게"),       # S4
])
def test_spc_alarms_in_scenario_month(grade, prop, line, start, end, direction):
    r = spc.monitor(A.filter_coils(grade, line), prop, start, end, A.spec(grade)[prop])
    a = r["first_alarm"]
    assert a and a["direction"] == direction and start <= a["date"] <= end


def test_spc_quiet_month_and_note():
    r = spc.monitor(A.filter_coils("HSLA590C", "CAL"), "YS", "2026-01-01", "2026-01-31", A.spec("HSLA590C")["YS"])
    assert r["first_alarm"] is None and r["baseline"]["n"] > 30
    assert "이상 신호가 없어요" in N.note_spc(r)[0]


def test_thickness_groups_and_defect_by_grade():
    tg = A.thickness_groups(A.filter_coils("HSLA590C"), "YS", A.spec("HSLA590C")["YS"])
    assert list(tg["band"]) == ["~0.75", "0.75~1.0", "1.0~1.4", "1.4~"] and tg["n"].sum() > 1000
    rows = A.defect_by_grade("파우더링", None, "2026-06-01", "2026-06-30")
    assert "합금화 온도" in set(rows["top_factor"]) and (rows["rate_target_pct"] >= 0).all()
    assert "BH340_OUT" in set(rows["grade"])


def test_sigma_changes_problem_coil_count():
    df = A.filter_coils("HSLA590C")
    n1 = len(A.comparison_groups(df, "HSLA590C", "YS", sigma=1.0)[2])
    n2 = len(A.comparison_groups(df, "HSLA590C", "YS", sigma=2.0)[2])
    assert n1 > n2 > 0  # 스펙 이탈 코일은 σ와 상관없이 남는다


def test_state_v4_and_v3_compat():
    st0 = ds.defaults()
    assert st0["prop"] == "TS" and st0["props"] == ["YS", "TS", "EL"] and ds.focus_prop(st0) == "TS" and st0["sigma"] == 1.0  # v5 기본 YP·TS·EL
    assert ds.shown_props(st0 | {"props": []}) == ["TS"]
    new, _, err = ds.apply_update(st0, {"props": ["YS", "EL"], "sigma": 1.34, "spec_cols": ["spec_roll", "spec_fin"],
                                        "specs": ["roll:CR590HSA", "fin:CR590HSC"]})
    assert not err and new["prop"] == "YS" and new["sigma"] == 1.3 and ds.focus_prop(new) == "YS"
    df = A.filter_coils(new["grade"], **ds.subset_args(new))
    assert ((df["spec_roll"] == "CR590HSA") | (df["spec_fin"] == "CR590HSC")).all() and len(df)
    old, _, _ = ds.apply_update(st0, {"spec_col": "spec_fin", "view": "코일 단건 분석"})  # v3 링크
    assert old["spec_cols"] == ["spec_fin"] and old["view"] == "코일분석"


def test_favorites_roundtrip(tmp_path, monkeypatch):
    from ui import favorites as F
    monkeypatch.setattr(F, "PATH", tmp_path / "favorites.json")
    st0 = ds.defaults()
    F.add("590C 전체", st0)
    F.add("두번째", {**st0, "line": "CAL"})
    items = F.load()
    assert [i["name"] for i in items] == ["두번째", "590C 전체"] and items[0]["state"]["line"] == "CAL"
    restored, _, err = ds.apply_update(ds.defaults(), items[0]["state"])
    assert not err and restored["line"] == "CAL"
    assert [i["name"] for i in F.remove("두번째")] == ["590C 전체"]


def _report_ctx():
    grade, prop = "HSLA590C", "YS"
    coils = A.filter_coils(grade, "CAL")
    target, baseline = A.split_period(coils, "2026-05-01", "2026-05-31")
    diag = A.diagnose(grade, "CAL", "2026-05-01", "2026-05-31")
    comp = A.compare(grade, prop, "CAL", "2026-05-01", "2026-05-31")
    sp = A.spec(grade)
    res = spc.monitor(coils, prop, "2026-05-01", "2026-05-31", sp[prop])
    tg = A.thickness_groups(coils, prop, sp[prop])
    notes = {"distribution": N.note_distribution(diag, prop, "이 기간 코일", "나머지 기간 코일"),
             "thickness": N.note_thickness(tg, prop), "spc": N.note_spc(res),
             "ranking": N.note_contribution(comp, "차이 크기"), "scatter": N.note_scatter(comp, "CT", prop)}
    return {"grade": grade, "focus_prop": prop, "headline": "재질 요약", "summary_lines": ["요약 1", "요약 2"],
            "summary": {q: {"spec": "-", "mean": "1", "std": "1", "rate": "-"} for q in ("YS", "TS")},
            "yearly": A.production(coils)[0], "dist_group": target, "dist_ref": baseline, "group": target,
            "ref": baseline, "group_name": "이 기간 코일", "ref_name": "나머지 기간 코일", "group_color": "#1E3A8A",
            "spec": sp, "coils": coils, "bands": A.thickness_band(coils["thk"]), "thickness": {prop: tg}, "spc": res,
            "comp": comp, "score_label": "차이 크기", "model_name": "차이 크기", "factor": "CT", "actions": ["CT 점검"],
            "notes": notes, "defect": "파우더링", "defect_grades": None, "defect_monthly": None,
            "cover_title": "HSLA590C 품질 분석 · 2026-05", "footer": "조건 · 테스트"}


def test_ppt_report_three_slides_inside_panels():
    from pptx import Presentation
    from tools import report_ppt as R
    ctx = _report_ctx()
    keys = ["summary", "distribution", "thickness", "spc", "ranking", "scatter"]
    assert R.check_count(len(keys)) is None and R.check_count(0) and R.check_count(10) is None  # 장수 제한 없음
    assert set(keys) <= set(R.available(ctx))
    pages = R.pages(keys, ctx)
    assert len(pages) == 3 and pages[0]["title"].startswith("HSLA590C 품질 분석")
    prs = Presentation(io.BytesIO(R.build_pptx(pages, ctx["footer"])))
    assert len(prs.slides) == 3
    for slide in prs.slides:
        title, subs, panels = R._layout_shapes(slide, prs.slide_height)
        assert title.text_frame.text and all(s.text_frame.text for s in subs)
        pics = [s for s in slide.shapes if s.shape_type == 13]
        assert len(pics) == 2
        for pic in pics:
            panel = min(panels, key=lambda p: abs(p.left - pic.left))
            assert panel.left <= pic.left and pic.left + pic.width <= panel.left + panel.width
            assert pic.top + pic.height <= panel.top + panel.height
    previews = R.preview_pngs(pages, ctx["footer"])
    assert len(previews) == 3 and all(p[:4] == b"\x89PNG" for p in previews)
